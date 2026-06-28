from abc import ABC, abstractmethod
import csv
import io
import logging
from html.parser import HTMLParser
from pathlib import Path

from app.exceptions import EmptyFileError, FileParseError, UnsupportedFileTypeError

logger = logging.getLogger(__name__)


class BaseParser(ABC):
    @abstractmethod
    def parse(self, data: bytes, filename: str) -> str: ...


class TextParser(BaseParser):
    def parse(self, data: bytes, filename: str) -> str:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        result = text.strip()
        if not result:
            raise EmptyFileError(filename)
        return result

class PDFParserExtended(BaseParser):
    """Layout-aware PDF parser for financial documents.

    Unlike :class:`PDFParser`, which uses pypdf's default text extraction (it
    reads characters in content-stream order and tends to collapse multi-column
    tables into an unreadable jumble), this parser uses pypdf's layout mode so
    columnar data — balance sheets, P&L statements, schedules — keeps its
    spatial alignment. That alignment is what lets the downstream LLM tell one
    column from the next. Pages are emitted with markers so retrieved chunks
    can be traced back to a page, and each page falls back to plain extraction
    if layout mode fails on it.
    """

    def parse(self, data: bytes, filename: str) -> str:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages: list[str] = []
            for index, page in enumerate(reader.pages, start=1):
                extracted = self._extract_page(page)
                if extracted:
                    pages.append(f"[Page {index}]\n{extracted}")
            result = "\n\n".join(pages).strip()
            logger.info("PDG Text Parsed %s", result)
        except Exception as exc:
            logger.warning("Failed to parse PDF %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc
        if not result:
            raise EmptyFileError(filename)
        return result

    @staticmethod
    def _extract_page(page) -> str:
        """Extract one page, preferring layout mode and falling back to plain."""
        try:
            text = page.extract_text(extraction_mode="layout")
        except Exception:  # noqa: BLE001 — layout mode unsupported/failed for this page
            text = page.extract_text()
        return (text or "").strip()

class PDFParser(BaseParser):
    def parse(self, data: bytes, filename: str) -> str:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages.append(extracted)
            result = "\n\n".join(pages).strip()
        except Exception as exc:
            logger.warning("Failed to parse PDF %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc
        if not result:
            raise EmptyFileError(filename)
        return result


class DocxParser(BaseParser):
    def parse(self, data: bytes, filename: str) -> str:
        try:
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts: list[str] = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = "\t".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        parts.append(row_text)
            result = "\n".join(parts).strip()
        except Exception as exc:
            logger.warning("Failed to parse DOCX %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc
        if not result:
            raise EmptyFileError(filename)
        return result


class ExcelParser(BaseParser):
    def parse(self, data: bytes, filename: str) -> str:
        ext = Path(filename).suffix.lower()
        try:
            if ext == ".xlsx":
                result = self._parse_xlsx(data)
            else:
                result = self._parse_xls(data)
        except (EmptyFileError, FileParseError):
            raise
        except Exception as exc:
            logger.warning("Failed to parse Excel %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc
        if not result.strip():
            raise EmptyFileError(filename)
        return result

    def _parse_xlsx(self, data: bytes) -> str:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections: list[str] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
            if rows:
                sections.append(f"## Sheet: {sheet_name}\n" + "\n".join(rows))
        return "\n\n".join(sections)

    def _parse_xls(self, data: bytes) -> str:
        import xlrd

        wb = xlrd.open_workbook(file_contents=data)
        sections: list[str] = []
        for sheet in wb.sheets():
            rows: list[str] = []
            for rx in range(sheet.nrows):
                cells = [str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols)]
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
            if rows:
                sections.append(f"## Sheet: {sheet.name}\n" + "\n".join(rows))
        return "\n\n".join(sections)


class CsvParser(BaseParser):
    def parse(self, data: bytes, filename: str) -> str:
        try:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
            reader = csv.reader(io.StringIO(text))
            rows = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
            result = "\n".join(rows).strip()
        except Exception as exc:
            logger.warning("Failed to parse CSV %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc
        if not result:
            raise EmptyFileError(filename)
        return result


class XmlParser(BaseParser):
    """Parses XBRL financial instances into a flat, LLM-efficient fact table.

    Uses py-xbrl (https://pypi.org/project/py-xbrl/) to extract structured facts
    (concept, value, unit, period) and renders them as a compact TSV instead of
    the verbose nested XML — far cheaper on tokens. Falls back to a generic
    element walk for plain (non-XBRL) XML.
    """

    # Cap a single (possibly huge) text fact so notes/narratives don't bloat the prompt.
    _MAX_VALUE_CHARS = 300

    def parse(self, data: bytes, filename: str) -> str:
        result = self._parse_xbrl(data, filename)
        if not result:
            result = self._parse_plain_xml(data, filename)
        if not result or not result.strip():
            raise EmptyFileError(filename)
        return result

    # --- XBRL via py-xbrl ---------------------------------------------------

    def _parse_xbrl(self, data: bytes, filename: str) -> str | None:
        try:
            from xbrl.cache import HttpCache
            from xbrl.instance import XbrlParser as PyXbrlParser
        except Exception as exc:  # noqa: BLE001 — py-xbrl optional; fall back to plain XML
            logger.warning("[XmlParser] py-xbrl unavailable (%s); using plain XML parser", exc)
            return None

        import os
        import tempfile

        cache_dir = os.environ.get("XBRL_CACHE_DIR") or os.path.join(
            tempfile.gettempdir(), "xbrl_cache"
        )
        tmp_path: str | None = None
        try:
            os.makedirs(cache_dir, exist_ok=True)
            # py-xbrl reads from a path/URL (it resolves the linked taxonomy), so
            # spill the in-memory bytes to a temp file first.
            suffix = Path(filename).suffix or ".xml"
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(data)
                tmp_path = tmp.name

            cache = HttpCache(cache_dir)
            try:  # some taxonomy hosts reject default clients
                cache.set_headers({"User-Agent": "ledger-lens-rag/1.0 (KPI workflow)"})
            except Exception:  # noqa: BLE001
                pass

            instance = PyXbrlParser(cache).parse_instance(tmp_path)
            flat = self._flatten(instance)
            if flat:
                logger.info(
                    "[XmlParser] parsed XBRL %s -> %d fact row(s)", filename, flat.count("\n")
                )
                return flat
            return None
        except Exception as exc:  # noqa: BLE001 — not XBRL / no network for taxonomy
            logger.warning(
                "[XmlParser] py-xbrl could not parse %s as XBRL (%s); falling back to plain XML",
                filename,
                exc,
            )
            return None
        finally:
            if tmp_path:
                try:
                    os.unlink(tmp_path)
                except OSError:
                    pass

    def _flatten(self, instance) -> str:
        """Render XBRL facts as a deduped `concept\tvalue\tunit\tperiod` table."""
        lines = ["concept\tvalue\tunit\tperiod"]
        seen: set[tuple] = set()
        for fact in getattr(instance, "facts", None) or []:
            concept = getattr(getattr(fact, "concept", None), "name", None) or ""
            value = getattr(fact, "value", None)
            if value is None:
                continue
            value_str = " ".join(str(value).split())
            if not value_str:
                continue
            if len(value_str) > self._MAX_VALUE_CHARS:
                value_str = value_str[: self._MAX_VALUE_CHARS] + "…"
            unit = self._format_unit(getattr(fact, "unit", None))
            period = self._format_period(getattr(fact, "context", None))
            key = (concept, value_str, unit, period)
            if key in seen:
                continue
            seen.add(key)
            lines.append(f"{concept}\t{value_str}\t{unit}\t{period}")
        return "\n".join(lines) if len(lines) > 1 else ""

    @staticmethod
    def _format_period(ctx) -> str:
        if ctx is None:
            return ""
        instant = getattr(ctx, "instant_date", None)
        if instant is not None:
            return str(instant)
        start = getattr(ctx, "start_date", None)
        end = getattr(ctx, "end_date", None)
        if start is not None or end is not None:
            return f"{start}/{end}"
        return "forever"

    @classmethod
    def _format_unit(cls, unit) -> str:
        if unit is None:
            return ""
        simple = getattr(unit, "unit", None)
        if simple:
            return cls._short(simple)
        numerator = getattr(unit, "numerator", None)
        denominator = getattr(unit, "denominator", None)
        if numerator or denominator:
            return f"{cls._short(numerator)}/{cls._short(denominator)}"
        return ""

    @staticmethod
    def _short(token) -> str:
        # "iso4217:INR" -> "INR", "xbrli:shares" -> "shares"
        return str(token).split(":")[-1] if token else ""

    # --- plain XML fallback -------------------------------------------------

    def _parse_plain_xml(self, data: bytes, filename: str) -> str:
        import xml.etree.ElementTree as ET

        try:
            root = ET.fromstring(data)
        except Exception as exc:
            logger.warning("Failed to parse XML %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc

        lines: list[str] = []

        def _walk(element) -> None:
            text = (element.text or "").strip()
            if text:
                lines.append(f"{element.tag}: {text}")
            for child in element:
                _walk(child)

        _walk(root)
        return "\n".join(lines).strip()


class _HtmlTextExtractor(HTMLParser):
    """Collects visible text from HTML, rendering tables as tab-separated rows.

    Financial filings published as HTML lean heavily on tables, so cell and row
    boundaries are preserved (cells joined by tabs, rows by newlines) to keep
    columnar data legible for the downstream LLM. ``<script>``/``<style>`` and
    other non-visible content is dropped.
    """

    _SKIP_TAGS = {"script", "style", "head", "meta", "link", "title"}
    _BLOCK_TAGS = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0
        self._cell_buffer: list[str] = []
        self._in_cell = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        elif tag in ("td", "th"):
            self._in_cell = True
            self._cell_buffer = []
        elif tag == "tr":
            self._parts.append("\n")
        elif tag in self._BLOCK_TAGS:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in ("td", "th"):
            self._parts.append(" ".join(self._cell_buffer).strip())
            self._parts.append("\t")
            self._in_cell = False
            self._cell_buffer = []

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = data.strip()
        if not text:
            return
        if self._in_cell:
            self._cell_buffer.append(text)
        else:
            self._parts.append(text + " ")

    def get_text(self) -> str:
        lines = "".join(self._parts).splitlines()
        # Trim trailing tab separators and drop blank lines.
        cleaned = [line.strip().rstrip("\t").strip() for line in lines]
        return "\n".join(line for line in cleaned if line)


class HtmlParser(BaseParser):
    def parse(self, data: bytes, filename: str) -> str:
        try:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
            extractor = _HtmlTextExtractor()
            extractor.feed(text)
            result = extractor.get_text().strip()
        except Exception as exc:
            logger.warning("Failed to parse HTML %s: %s", filename, exc)
            raise FileParseError(filename, str(exc)) from exc
        if not result:
            raise EmptyFileError(filename)
        return result


_REGISTRY: dict[str, BaseParser] = {
    ".pdf": PDFParser(),
    ".docx": DocxParser(),
    ".xlsx": ExcelParser(),
    ".xls": ExcelParser(),
    ".txt": TextParser(),
    ".csv": CsvParser(),
    ".md": TextParser(),
    ".xml": XmlParser(),
    ".xbrl": XmlParser(),
    ".html": HtmlParser(),
    ".htm": HtmlParser(),
}


def get_parser(filename: str) -> BaseParser:
    logger.info("File name to be parsed: %s", filename)
    ext = Path(filename).suffix.lower()
    if ext not in _REGISTRY:
        logger.warning("Unsupported file type %r for file %s", ext, filename)
        raise UnsupportedFileTypeError(ext)
    parser = _REGISTRY[ext]
    logger.debug("Selected %s for %s", type(parser).__name__, filename)
    return parser


def supported_extensions() -> list[str]:
    return list(_REGISTRY.keys())
