"""Unit tests for app/core/parsers.py — all using in-memory bytes."""

import csv
import io

import pytest

from app.exceptions import EmptyFileError, FileParseError, UnsupportedFileTypeError
from app.core.parsers import (
    CsvParser,
    DocxParser,
    ExcelParser,
    PDFParser,
    TextParser,
    XmlParser,
    get_parser,
)


# ---------------------------------------------------------------------------
# Helpers to build minimal valid in-memory file bytes
# ---------------------------------------------------------------------------

def _make_pdf(text: str) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import NameObject

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    page = writer.pages[0]
    # Inject a simple content stream with the text so extract_text returns it
    content = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    from pypdf.generic import ArrayObject, DictionaryObject, DecodedStreamObject, NumberObject
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = writer._add_object(stream)
    # Register a minimal font so PDF is well-formed
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    resources = DictionaryObject({
        NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): writer._add_object(font),
        })
    })
    page[NameObject("/Resources")] = resources
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(sheet_name: str, rows: list[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# TextParser
# ---------------------------------------------------------------------------

class TestTextParser:
    parser = TextParser()

    def test_parses_utf8(self):
        data = b"Hello, world!"
        assert "Hello" in self.parser.parse(data, "file.txt")

    def test_parses_latin1_fallback(self):
        data = "caf\xe9".encode("latin-1")
        result = self.parser.parse(data, "file.txt")
        assert "caf" in result

    def test_empty_raises(self):
        with pytest.raises(EmptyFileError):
            self.parser.parse(b"   \n  ", "file.txt")

    def test_zero_bytes_raises(self):
        with pytest.raises(EmptyFileError):
            self.parser.parse(b"", "file.txt")


# ---------------------------------------------------------------------------
# PDFParser
# ---------------------------------------------------------------------------

class TestPDFParser:
    parser = PDFParser()

    def test_extracts_text(self):
        data = _make_pdf("Invoice total 500")
        result = self.parser.parse(data, "invoice.pdf")
        assert "Invoice" in result or "500" in result

    def test_corrupt_raises_file_parse_error(self):
        with pytest.raises(FileParseError):
            self.parser.parse(b"not a pdf at all", "bad.pdf")

    def test_empty_bytes_raises_file_parse_error(self):
        with pytest.raises((FileParseError, EmptyFileError)):
            self.parser.parse(b"", "empty.pdf")


# ---------------------------------------------------------------------------
# DocxParser
# ---------------------------------------------------------------------------

class TestDocxParser:
    parser = DocxParser()

    def test_extracts_paragraphs(self):
        data = _make_docx("Financial report Q1")
        result = self.parser.parse(data, "report.docx")
        assert "Financial report Q1" in result

    def test_corrupt_raises_file_parse_error(self):
        with pytest.raises(FileParseError):
            self.parser.parse(b"not a docx", "bad.docx")

    def test_empty_document_raises_empty_file_error(self):
        from docx import Document
        doc = Document()
        buf = io.BytesIO()
        doc.save(buf)
        with pytest.raises(EmptyFileError):
            self.parser.parse(buf.getvalue(), "empty.docx")


# ---------------------------------------------------------------------------
# ExcelParser
# ---------------------------------------------------------------------------

class TestExcelParser:
    parser = ExcelParser()

    def test_xlsx_extracts_rows(self):
        data = _make_xlsx("Sheet1", [["Name", "Amount"], ["Alice", 100]])
        result = self.parser.parse(data, "data.xlsx")
        assert "Name" in result
        assert "Alice" in result

    def test_xlsx_includes_sheet_header(self):
        data = _make_xlsx("Ledger", [["Account", "Balance"]])
        result = self.parser.parse(data, "ledger.xlsx")
        assert "Ledger" in result

    def test_xlsx_corrupt_raises_file_parse_error(self):
        with pytest.raises(FileParseError):
            self.parser.parse(b"not an xlsx", "bad.xlsx")

    def test_xlsx_empty_workbook_raises_empty_file_error(self):
        import openpyxl
        wb = openpyxl.Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(EmptyFileError):
            self.parser.parse(buf.getvalue(), "empty.xlsx")


# ---------------------------------------------------------------------------
# CsvParser
# ---------------------------------------------------------------------------

class TestCsvParser:
    parser = CsvParser()

    def test_parses_rows(self):
        data = _make_csv([["Date", "Amount"], ["2024-01-01", "500"]])
        result = self.parser.parse(data, "data.csv")
        assert "Date" in result
        assert "500" in result

    def test_empty_raises(self):
        with pytest.raises(EmptyFileError):
            self.parser.parse(b"", "empty.csv")

    def test_corrupt_bytes_latin1_fallback(self):
        data = b"\xff\xfe" + b"col1,col2\nval1,val2"
        result = self.parser.parse(data, "data.csv")
        assert result  # should not raise


# ---------------------------------------------------------------------------
# XmlParser
# ---------------------------------------------------------------------------

class TestXmlParser:
    parser = XmlParser()

    def test_extracts_text(self):
        data = b"<invoice><total>500</total><customer>Alice</customer></invoice>"
        result = self.parser.parse(data, "invoice.xml")
        assert "500" in result
        assert "Alice" in result

    def test_malformed_raises_file_parse_error(self):
        with pytest.raises(FileParseError):
            self.parser.parse(b"<unclosed>", "bad.xml")

    def test_empty_elements_raise_empty_file_error(self):
        with pytest.raises(EmptyFileError):
            self.parser.parse(b"<root><child/></root>", "empty.xml")


# ---------------------------------------------------------------------------
# get_parser registry
# ---------------------------------------------------------------------------

class TestGetParser:
    def test_known_extensions_resolve(self):
        for ext in (".pdf", ".docx", ".xlsx", ".xls", ".txt", ".csv", ".md", ".xml"):
            parser = get_parser(f"file{ext}")
            assert parser is not None

    def test_unknown_extension_raises(self):
        with pytest.raises(UnsupportedFileTypeError):
            get_parser("file.pptx")

    def test_no_extension_raises(self):
        with pytest.raises(UnsupportedFileTypeError):
            get_parser("noextension")

    def test_doc_legacy_raises(self):
        with pytest.raises(UnsupportedFileTypeError):
            get_parser("old.doc")
